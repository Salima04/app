"""Document processing + RAG retrieval."""
import io
import re
import hashlib
import logging
from typing import List, Dict, Tuple
from pypdf import PdfReader
from docx import Document as DocxDocument
import openpyxl
import csv

from llm_service import embed_text, cosine

log = logging.getLogger("vidyagpt.rag")


def extract_text(filename: str, content: bytes) -> List[Tuple[int, str]]:
    """Return list of (page_number, text) tuples."""
    lower = filename.lower()
    pages: List[Tuple[int, str]] = []
    try:
        if lower.endswith(".pdf"):
            reader = PdfReader(io.BytesIO(content))
            for i, pg in enumerate(reader.pages, start=1):
                t = (pg.extract_text() or "").strip()
                if t:
                    pages.append((i, t))
        elif lower.endswith(".docx"):
            doc = DocxDocument(io.BytesIO(content))
            buf = []
            for p in doc.paragraphs:
                if p.text.strip():
                    buf.append(p.text)
            for tb in doc.tables:
                for row in tb.rows:
                    buf.append(" | ".join(c.text for c in row.cells))
            pages.append((1, "\n".join(buf)))
        elif lower.endswith(".xlsx"):
            wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
            for si, ws in enumerate(wb.worksheets, start=1):
                rows = []
                for row in ws.iter_rows(values_only=True):
                    rows.append(" | ".join(str(c) if c is not None else "" for c in row))
                pages.append((si, f"[Sheet: {ws.title}]\n" + "\n".join(rows)))
        elif lower.endswith(".csv"):
            reader = csv.reader(io.StringIO(content.decode("utf-8", errors="ignore")))
            rows = [" | ".join(r) for r in reader]
            pages.append((1, "\n".join(rows)))
        else:  # txt and others
            pages.append((1, content.decode("utf-8", errors="ignore")))
    except Exception as e:
        log.exception("document extraction failed")
        pages.append((1, "[Unable to extract text from this file]"))
    return pages


def chunk_text(text: str, chunk_size: int = 800, overlap: int = 120) -> List[str]:
    text = re.sub(r"\s+", " ", text).strip()
    if not text:
        return []
    chunks = []
    start = 0
    while start < len(text):
        end = min(len(text), start + chunk_size)
        # try to end at sentence boundary
        if end < len(text):
            dot = text.rfind(". ", start + chunk_size - 200, end)
            if dot > start + 200:
                end = dot + 1
        chunks.append(text[start:end].strip())
        start = end - overlap
        if start < 0:
            start = 0
        if end == len(text):
            break
    return [c for c in chunks if len(c) > 30]


def build_chunks(pages: List[Tuple[int, str]]) -> List[Dict]:
    out = []
    idx = 0
    for pg, text in pages:
        for ch in chunk_text(text):
            out.append({
                "chunk_index": idx,
                "page": pg,
                "text": ch,
                "embedding": embed_text(ch),
            })
            idx += 1
    return out


def retrieve(query: str, chunks: List[Dict], top_k: int = 5, min_score: float = 0.02) -> List[Dict]:
    """Hybrid retrieval: cosine over hash-TF embeddings + lexical overlap fallback.
    Guarantees top matches when any keyword overlap exists.
    """
    if not chunks:
        return []
    q_emb = embed_text(query)
    q_tokens = set(re.findall(r"[a-zA-Z\u0900-\u097F0-9]{3,}", query.lower()))
    scored = []
    for c in chunks:
        emb = c.get("embedding") or []
        sem = cosine(q_emb, emb) if emb else 0.0
        # lexical overlap boost
        c_tokens = set(re.findall(r"[a-zA-Z\u0900-\u097F0-9]{3,}", c.get("text", "").lower()))
        overlap = (len(q_tokens & c_tokens) / max(1, len(q_tokens))) if q_tokens else 0.0
        combined = 0.6 * sem + 0.4 * overlap
        scored.append((combined, c))
    scored.sort(key=lambda x: x[0], reverse=True)
    # first pass: threshold
    top = [{**c, "score": round(float(s), 4)} for s, c in scored[:top_k] if s >= min_score]
    # lexical fallback: if none pass threshold but we have keyword overlap somewhere
    if not top and q_tokens:
        for s, c in scored[:top_k]:
            c_tokens = set(re.findall(r"[a-zA-Z\u0900-\u097F0-9]{3,}", c.get("text", "").lower()))
            if q_tokens & c_tokens:
                top.append({**c, "score": round(float(s), 4)})
        top = top[:top_k]
    return top


def normalize_question(q: str) -> str:
    """Normalize for cache lookup."""
    q = q.lower().strip()
    q = re.sub(r"[^\w\u0900-\u097F\s]", " ", q)
    q = re.sub(r"\s+", " ", q)
    # remove common stopwords
    stops = {"a", "an", "the", "is", "are", "of", "in", "on", "for", "to", "and", "please", "kya", "kaise", "hai", "ka", "ki", "ke"}
    q = " ".join(w for w in q.split() if w not in stops)
    return q


def question_hash(q_norm: str) -> str:
    return hashlib.sha256(q_norm.encode()).hexdigest()
